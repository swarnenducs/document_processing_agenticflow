"""Tests for style extraction, mapping, generation, validation, and tools."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document

from document_processing_agenticflow.graph import build_graph
from document_processing_agenticflow.services.confidence import build_confidence_report
from document_processing_agenticflow.services.document_generator import generate_styled_document
from document_processing_agenticflow.services.document_validator import validate_documents
from document_processing_agenticflow.services.field_mapper import map_json_to_template
from document_processing_agenticflow.services.style_extractor import extract_word_styles
from document_processing_agenticflow.tools import get_document_tools
from scripts.create_sample_template import build_sample_template


SAMPLE_DATA = {
    "invoice_number": "INV-99",
    "invoice_date": "2026-07-10",
    "customer": {
        "name": "Test Co",
        "address": "1 Test Rd",
        "email": "t@example.com",
    },
    "items": [
        {"description": "Item A", "quantity": 1, "unit_price": 10, "total": 10},
        {"description": "Item B", "quantity": 2, "unit_price": 5, "total": 10},
    ],
    "subtotal": 20,
    "tax": 4,
    "total_amount": 24,
    "notes": "Paid",
}


def test_extract_styles_and_placeholders(tmp_path: Path) -> None:
    template = build_sample_template(tmp_path / "template.docx")
    extracted = extract_word_styles(template)

    assert extracted.styles
    assert "invoice_number" in extracted.placeholders
    assert "customer.name" in extracted.placeholders
    assert "total_amount" in extracted.placeholders


def test_rule_based_mapping_has_confidence(tmp_path: Path) -> None:
    template = build_sample_template(tmp_path / "template.docx")
    extracted = extract_word_styles(template)
    mapping = map_json_to_template(extracted, SAMPLE_DATA)

    assert mapping.mapper_source == "rules"
    assert mapping.mapper_provider == "rules"
    assert mapping.mapping_confidence > 0
    assert mapping.coverage_score == 1.0
    inv = next(m for m in mapping.mappings if m.placeholder == "invoice_number")
    assert inv.value == "INV-99"
    assert 0 < inv.confidence <= 1


def test_generation_integrity_and_validation(tmp_path: Path) -> None:
    template = build_sample_template(tmp_path / "template.docx")
    extracted = extract_word_styles(template)
    mapping = map_json_to_template(extracted, SAMPLE_DATA)
    output = tmp_path / "out.docx"
    generation = generate_styled_document(extracted, mapping, output)

    assert generation.integrity_score >= 0.8
    assert generation.generation_confidence > 0
    assert not generation.leftover_placeholders

    validation = validate_documents(
        extracted, output, SAMPLE_DATA, mapping, prefer_llm=False
    )
    assert validation.passed
    assert validation.validation_score >= 0.7
    assert validation.validator_source == "rules"
    assert validation.validator_provider == "rules"

    report = build_confidence_report(mapping, generation, validation)
    assert report.overall_confidence >= 0.7
    assert report.per_field


def test_end_to_end_graph_with_validation(tmp_path: Path) -> None:
    template = build_sample_template(tmp_path / "template.docx")
    data_path = tmp_path / "data.json"
    data_path.write_text(json.dumps(SAMPLE_DATA), encoding="utf-8")
    output = tmp_path / "out.docx"

    result = build_graph().invoke(
        {
            "template_path": str(template),
            "data_path": str(data_path),
            "output_path": str(output),
            "errors": [],
            "status": "started",
            "retry_count": 0,
            "max_retries": 1,
            "validation_threshold": 0.7,
            "skip_validation": False,
        }
    )

    assert result["status"] == "completed"
    assert output.exists()
    assert result.get("validation") is not None
    assert result["validation"].passed
    assert result.get("confidence") is not None
    assert result["confidence"].overall_confidence > 0

    doc = Document(output)
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "INV-99" in full_text
    assert "Test Co" in full_text
    assert "{{invoice_number}}" not in full_text


def test_tools_are_registered() -> None:
    tools = get_document_tools()
    names = {t.name for t in tools}
    assert names == {
        "load_json_data",
        "extract_word_styles",
        "map_json_to_template",
        "generate_styled_document",
        "validate_documents",
        "compute_confidence_report",
    }
