"""Detect fill tokens, stamp LLM marker plans, closest library match."""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document

from document_processing_mcp.services.library_match import (
    find_closest_library_template,
    similarity_score,
)
from document_processing_mcp.services.marker_apply import apply_marker_plan, normalize_marker_token
from document_processing_mcp.services.placeholders import template_has_markers
from document_processing_mcp.services.prompts.loader import load_prompt_yaml
from document_processing_mcp.services.style_extractor import extract_word_styles
from doc_sample_template import build_sample_template


def test_invoice_template_has_markers(tmp_path: Path) -> None:
    path = build_sample_template(tmp_path / "inv.docx")
    extracted = extract_word_styles(path)
    has, keys = template_has_markers(extracted)
    assert has
    assert "invoice_number" in keys or any("invoice" in k.lower() for k in keys)


def test_plain_docx_has_no_markers(tmp_path: Path) -> None:
    path = tmp_path / "plain.docx"
    doc = Document()
    doc.add_paragraph("Customer Name: Acme Corporation")
    doc.add_paragraph("Effective Date: August 28, 2026")
    doc.save(path)
    extracted = extract_word_styles(path)
    has, keys = template_has_markers(extracted, template_path=str(path))
    assert not has
    assert keys == []


def test_apply_marker_plan_replaces_sample_text(tmp_path: Path) -> None:
    src = tmp_path / "plain.docx"
    doc = Document()
    doc.add_paragraph("Customer Name: Acme Corporation")
    doc.save(src)
    plan = SimpleNamespace(
        markers=[
            SimpleNamespace(
                sample_value="Acme Corporation",
                marker="<customer_name>",
                confidence="HIGH",
            )
        ],
        tables=[],
    )
    dest = tmp_path / "marked.docx"
    result = apply_marker_plan(src, dest, plan)
    assert result["markers_present"]
    extracted = extract_word_styles(dest)
    assert "customer_name" in extracted.placeholders
    text = "\n".join(b.text for b in extracted.blocks)
    assert "<customer_name>" in text
    assert "Acme Corporation" not in text


def test_normalize_marker_token() -> None:
    assert normalize_marker_token("Customer Name") == "<customer_name>"
    assert normalize_marker_token("<already_ok>") == "<already_ok>"


def test_similarity_prefers_overlapping_prose() -> None:
    assert similarity_score("product pricing catalog items", "product pricing list") > 0.2
    assert similarity_score("aaa", "zzz") == 0.0


def test_closest_library_skips_self(tmp_path: Path, monkeypatch) -> None:
    from document_processing_mcp.services import library_match

    monkeypatch.setattr(library_match, "library_template_dirs", lambda: [tmp_path])
    a = tmp_path / "a.docx"
    b = tmp_path / "b.docx"
    da, db = Document(), Document()
    da.add_paragraph("GPO agreement product pricing catalog items")
    db.add_paragraph("GPO agreement product pricing catalog services")
    da.save(a)
    db.save(b)
    extracted = extract_word_styles(a)
    match = find_closest_library_template(extracted, exclude_path=a)
    assert match is not None
    assert Path(match["path"]).name == "b.docx"


def test_marker_synthesizer_prompt_versioned() -> None:
    payload = load_prompt_yaml("marker_synthesizer.yml")
    assert payload["version"] == "1.0.0"
    assert payload["required_version"] == "1.0.0"
    assert "snake_case" in payload["system"]
    assert "{document_text}" in payload["human"]
    assert "{reference_placeholders_json}" in payload["human"]


def test_unmarked_template_fails_when_synthesis_disabled(
    tmp_path: Path, monkeypatch
) -> None:
    from document_processing_mcp.core.settings import reload_settings
    from document_processing_mcp.nodes.pipeline import synthesize_markers_node

    monkeypatch.setenv("DOCUMENT_MARKER_SYNTHESIS_ENABLED", "false")
    reload_settings()
    path = tmp_path / "plain.docx"
    doc = Document()
    doc.add_paragraph("Customer Name: Acme Corporation")
    doc.save(path)
    extracted = extract_word_styles(path)
    result = synthesize_markers_node(
        {
            "extracted": extracted,
            "json_data": {"customer": "Acme"},
            "output_path": str(tmp_path / "out.docx"),
            "status": "styles_extracted",
            "errors": [],
        }
    )
    assert result["status"] == "failed"
    assert result["marker_detection"]["reason"] == "marker_synthesis_disabled"
    assert any("DOCUMENT_MARKER_SYNTHESIS_ENABLED is off" in e for e in result["errors"])


def test_tagged_template_skips_synthesis_when_disabled(
    tmp_path: Path, monkeypatch
) -> None:
    from document_processing_mcp.core.settings import reload_settings
    from document_processing_mcp.nodes.pipeline import synthesize_markers_node

    monkeypatch.setenv("DOCUMENT_NON_TAG_ENABLED", "false")
    monkeypatch.delenv("DOCUMENT_MARKER_SYNTHESIS_ENABLED", raising=False)
    reload_settings()
    path = build_sample_template(tmp_path / "inv.docx")
    extracted = extract_word_styles(path)
    result = synthesize_markers_node(
        {
            "extracted": extracted,
            "json_data": {"invoice_number": "INV-1"},
            "output_path": str(tmp_path / "out.docx"),
            "status": "styles_extracted",
            "errors": [],
        }
    )
    assert result["status"] == "markers_ready"
    assert result["marker_detection"]["reason"] == "marker_synthesis_disabled"
    assert result["marker_detection"]["had_markers"] is True


def test_synthesize_markers_if_needed_respects_env(tmp_path: Path, monkeypatch) -> None:
    from document_processing_mcp.core.settings import reload_settings
    from document_processing_mcp.services.marker_synthesizer import (
        synthesize_markers_if_needed,
    )

    monkeypatch.setenv("DOCUMENT_MARKER_SYNTHESIS_ENABLED", "false")
    reload_settings()
    path = tmp_path / "plain.docx"
    doc = Document()
    doc.add_paragraph("No tags here")
    doc.save(path)
    extracted = extract_word_styles(path)
    with pytest.raises(RuntimeError, match="DOCUMENT_MARKER_SYNTHESIS_ENABLED is off"):
        synthesize_markers_if_needed(extracted, {"x": 1})
