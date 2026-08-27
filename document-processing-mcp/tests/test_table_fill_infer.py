"""Infer product-table fills from headers + JSON arrays."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document

from document_processing_mcp.models.schemas import MappingResult
from document_processing_mcp.services.document_generator import generate_styled_document
from document_processing_mcp.services.style_extractor import extract_word_styles
from document_processing_mcp.services.table_fill_infer import infer_table_fills, match_header_to_field


def _empty_price_list_docx(path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("PRICE LIST")
    table = doc.add_table(rows=5, cols=6)
    headers = [
        "PRODUCT CODE",
        "PRODUCT DESCRIPTION",
        "EA/UOM",
        "UOM",
        "PRICE/EA",
        "PRICE/UOM",
    ]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    doc.save(path)
    return path


def test_match_gpo_price_headers() -> None:
    keys = ["productCode", "productDescription", "eaUom", "uom", "eaPrice", "marketPrice"]
    assert match_header_to_field("PRODUCT CODE", keys) == "productCode"
    assert match_header_to_field("PRODUCT DESCRIPTION", keys) == "productDescription"
    assert match_header_to_field("UOM", keys) == "uom"
    assert match_header_to_field("PRICE/EA", keys) == "eaPrice"
    assert match_header_to_field("PRICE/UOM", keys) == "marketPrice"


def test_infer_and_generate_fills_empty_price_table(tmp_path: Path) -> None:
    template = _empty_price_list_docx(tmp_path / "price.docx")
    extracted = extract_word_styles(template)
    repo = Path(__file__).resolve().parents[2]
    data = json.loads((repo / "samples" / "data" / "gpo_agreement.json").read_text(encoding="utf-8"))
    fills = infer_table_fills(extracted, data)
    assert fills
    assert fills[0].array_json_path == "products"
    fields = {c.json_field for c in fills[0].columns}
    assert "productCode" in fields
    assert "productDescription" in fields

    mapping = MappingResult(table_fills=fills, mapping_confidence=1.0)
    output = tmp_path / "out.docx"
    generate_styled_document(extracted, mapping, output, json_data=data)

    filled = extract_word_styles(output)
    text = "\n".join(b.text for b in filled.blocks)
    assert "ABC-1001" in text
    assert "Industrial Grade Steel Bolt 10mm" in text
    assert "ABC-1002" in text
